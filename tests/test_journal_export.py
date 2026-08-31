# tests/test_journal_export.py — 筛选、导出与整表改删（v0.3.0 设计 §5.2 / §5.5）
#
# 这一层有三件事必须用**字节/往返**断言，而不是"看着没乱码"：
#
# 1. **CSV 的 BOM**（设计 §5.5 点名的那个坑）：没有它，Excel 按 GBK 解，
#    理由与名称全是乱码。而"乱码"在 pytest 里长得跟正常一模一样——
#    `df.to_csv()` 之后再 `read_csv()` 照样对得上。所以只能断言开头三个字节。
# 2. **前导零**：`000333` 在 Excel 里被当数字就变成 333。CSV 那边靠 dtype 守，
#    xlsx 那边靠"写进去的是文本单元格"，两条路都得用读回来的**类型**断言。
# 3. **导出 == 当前筛选结果**（所见即所得）：筛完再导，导出的行必须与页面上
#    看到的那些逐个 trade_id 对上。这条是设计对导出的唯一语义要求。
#
# 改删走 trade_id 而不是行号：编辑区显示的是**筛选后**的表，行号与整本日志对不上；
# 按行号写盘的话，用户改第 2 行会改掉别人那一行，而且不会报错。
import io
import sys
from datetime import date, datetime

import pandas as pd
import pytest

from quant.config import Costs, StampTaxRule
from quant.journal import export, schema, store

# 与 test_journal_pnl.py 同一套成本参数（费用一律走 backtest/costs.py）。
CFG = Costs(
    commission_rate=0.00025, commission_min=5.0, slippage=0.001,
    stamp_tax=(StampTaxRule(rate=0.001, until=date(2023, 8, 27)),
               StampTaxRule(rate=0.0005, frm=date(2023, 8, 28))),
)


def _frame(*rows) -> pd.DataFrame:
    """把若干 dict 补全默认值并编上可读的 trade_id，得到一张与线上同形的表。"""
    records = []
    for i, row in enumerate(rows, 1):
        record = schema.apply_defaults(row, costs=CFG)
        record["trade_id"] = record["trade_id"] or f"T{i}"
        records.append(record)
    # 走 store 的 dtype 归一：筛选/导出拿到的永远是 load_trades 出来的那种表，
    # 直接 DataFrame(records) 的 date 列是 object，测试就测不到真实形态。
    return store._coerce(pd.DataFrame(records, columns=list(schema.COLUMNS)))


def _sample() -> pd.DataFrame:
    """一张覆盖四种 kind、两只票、两个来源的小日志。"""
    return _frame(
        dict(date=date(2026, 8, 3), symbol="000333", name="美的集团", kind="buy",
             shares=1000.0, price=71.5, source="ma_cross", reason="20日线金叉"),
        dict(date=date(2026, 8, 10), symbol="600519", name="贵州茅台", kind="buy",
             shares=100.0, price=1600.0, source="discretionary", reason="跌到心理价位"),
        dict(date=date(2026, 8, 17), symbol="000333", kind="dividend", amount=320.0,
             source="ma_cross", reason="年度分红到账", note="含税"),
        dict(date=date(2026, 8, 24), symbol="000333", kind="sell",
             shares=1000.0, price=75.0, source="ma_cross", reason="跌破20日线"),
        dict(date=date(2026, 8, 27), symbol="600519", kind="adjust", shares=10.0,
             source="discretionary", reason="10送1", note="送股"),
    )


def ids(df) -> list[str]:
    return list(df["trade_id"])


# ================================================================ 筛选（设计 §5.2）

def test_no_filter_returns_everything():
    """全部条件留空 = 不筛。这是页面第一次打开时的状态：默认必须是"全都给我看"，
    而不是"什么都不给看"。"""
    df = _sample()
    assert ids(export.filter_trades(df)) == ids(df)


def test_empty_selections_mean_no_filter_not_nothing():
    """多选框空着 = 不限制这一项，**不是**"筛掉全部"。

    写成 `df[df.symbol.isin(symbols)]` 而不判空的话，一打开页面就是一张空表，
    用户只会以为日志丢了——而文件其实一个字节都没变。
    """
    df = _sample()
    out = export.filter_trades(df, symbols=[], kinds=[], sources=[], reason="")
    assert ids(out) == ids(df)


def test_date_range_is_inclusive_on_both_ends():
    """区间两端都算在内。差一天的话，用户选"8月3日到8月24日"会漏掉首尾两笔，
    而汇总数字照样给得出来。"""
    out = export.filter_trades(_sample(), start=date(2026, 8, 3), end=date(2026, 8, 24))
    assert ids(out) == ["T1", "T2", "T3", "T4"]


def test_only_start_or_only_end_is_allowed():
    df = _sample()
    assert ids(export.filter_trades(df, start=date(2026, 8, 17))) == ["T3", "T4", "T5"]
    assert ids(export.filter_trades(df, end=date(2026, 8, 10))) == ["T1", "T2"]


def test_start_after_end_gives_nothing_not_an_exception():
    """日期选反了是常事（两个独立的输入框）。给空表 + 页面自己说"没有符合的记录"，
    比抛异常把整页打没好——而且不能悄悄把两头调换：那是替用户改需求。"""
    out = export.filter_trades(_sample(), start=date(2026, 8, 24), end=date(2026, 8, 3))
    assert ids(out) == []


def test_symbols_filter_accepts_several_codes():
    out = export.filter_trades(_sample(), symbols=["600519"])
    assert ids(out) == ["T2", "T5"]
    assert set(out["symbol"]) == {"600519"}


def test_symbol_filter_matches_as_text_so_leading_zeros_survive():
    """`000333` 不许在筛选路径上被当成数字。这里筛的是字符串，
    传 333 应该什么都不匹配（而不是匹配上 000333）。"""
    df = _sample()
    assert ids(export.filter_trades(df, symbols=["000333"])) == ["T1", "T3", "T4"]
    assert ids(export.filter_trades(df, symbols=["333"])) == []


def test_kind_filter_separates_the_four_record_types():
    df = _sample()
    assert ids(export.filter_trades(df, kinds=["buy"])) == ["T1", "T2"]
    assert ids(export.filter_trades(df, kinds=["sell", "dividend"])) == ["T3", "T4"]
    assert ids(export.filter_trades(df, kinds=["adjust"])) == ["T5"]


def test_source_filter_is_the_point_of_the_whole_field():
    """按来源筛是 source 字段存在的理由：把"照信号做的"与"自己拍脑袋做的"分开看。"""
    out = export.filter_trades(_sample(), sources=["discretionary"])
    assert ids(out) == ["T2", "T5"]


def test_reason_keyword_is_a_substring_match():
    out = export.filter_trades(_sample(), reason="20日线")
    assert ids(out) == ["T1", "T4"]


def test_reason_keyword_ignores_case_and_surrounding_spaces():
    """关键词从文本框来，前后空格是手滑的常态；英文缩写（MA / ma）大小写也不该分家。"""
    df = _frame(dict(date=date(2026, 8, 3), symbol="000333", kind="buy", shares=100.0,
                     price=10.0, reason="MA金叉"))
    assert ids(export.filter_trades(df, reason="  ma  ")) == ["T1"]


def test_reason_keyword_does_not_leak_into_the_note_column():
    """筛的是**理由**。顺手把备注也搜进去的话，用户按"含税"筛出来一笔理由完全无关的
    记录，会以为筛选坏了；而两个字段分开是设计 §2.3 定的。"""
    assert ids(export.filter_trades(_sample(), reason="含税")) == []


def test_conditions_are_combined_with_and():
    out = export.filter_trades(_sample(), symbols=["000333"], kinds=["sell"],
                               start=date(2026, 8, 1))
    assert ids(out) == ["T4"]


def test_filter_keeps_columns_dtypes_and_order():
    """筛完的表还要喂给 compute_pnl 与导出。列少一个、dtype 变一次，
    后面全是"看着正常的错数字"。"""
    out = export.filter_trades(_sample(), symbols=["000333"])
    assert list(out.columns) == list(schema.COLUMNS)
    assert dict(out.dtypes.astype(str)) == dict(store.empty_trades().dtypes.astype(str))
    assert out["symbol"].tolist()[0] == "000333"
    assert pd.api.types.is_datetime64_any_dtype(out["date"])
    assert out["shares"].dtype == "float64"


def test_filter_does_not_touch_the_input_frame():
    """入参是整本日志（页面上还要接着用）。原地改一次，页面下半截看到的就是筛过的表。"""
    df = _sample()
    before = df.copy(deep=True)
    export.filter_trades(df, symbols=["600519"], reason="茅台")
    pd.testing.assert_frame_equal(df, before)


def test_filter_of_an_empty_journal_stays_an_empty_typed_frame():
    out = export.filter_trades(store.empty_trades(), symbols=["000333"])
    assert out.empty
    assert list(out.columns) == list(schema.COLUMNS)
    assert pd.api.types.is_datetime64_any_dtype(out["date"])


def test_filter_index_is_reset_so_row_numbers_mean_what_they_look_like():
    """筛完必须重排索引：页面上的编辑区按位置回读那张表，
    留着 1/3/4 这种洞的索引会让"第 0 行"指向别人。"""
    out = export.filter_trades(_sample(), symbols=["000333"])
    assert list(out.index) == [0, 1, 2]


# ================================================================ CSV 导出（§5.5）

def test_csv_starts_with_a_utf8_bom():
    """**字节断言**。设计 §5.5 点名的那个坑：没有 BOM，Excel 按 GBK 解中文全是乱码。
    而 pandas 读回来时 BOM 有没有都对，所以只能在字节层面守。"""
    blob = export.to_csv_bytes(_sample())
    assert blob[:3] == b"\xef\xbb\xbf", f"CSV 开头不是 UTF-8 BOM: {blob[:8]!r}"


def test_csv_keeps_chinese_readable_after_the_bom():
    blob = export.to_csv_bytes(_sample())
    text = blob.decode("utf-8-sig")
    assert "美的集团" in text and "20日线金叉" in text
    assert not text.startswith("﻿"), "utf-8-sig 解码后不该再留一个裸 BOM 字符"


def test_csv_uses_unix_line_endings_like_the_journal_file():
    """与 store.save_trades 同一个口径（换行符随平台变的话，一次导出的 diff 是整个文件）。"""
    blob = export.to_csv_bytes(_sample())
    assert b"\r\n" not in blob


def test_csv_contains_exactly_the_filtered_rows():
    """所见即所得：导出的是**当前筛选后**的结果，不是整本日志。"""
    filtered = export.filter_trades(_sample(), symbols=["000333"], kinds=["buy", "sell"])
    text = export.to_csv_bytes(filtered).decode("utf-8-sig")
    body = [line for line in text.strip().splitlines()[1:] if line]
    assert len(body) == 2, text
    assert "600519" not in text
    assert all(tid in text for tid in ids(filtered))


def test_csv_round_trips_through_the_journal_loader(tmp_path):
    """导出的 CSV 与 journal/trades.csv **同一种格式**：写回去能直接被 load_trades 读。

    这条不是锦上添花：用户误删一段记录时，最快的自救就是把导出的那份放回去。
    格式一旦分叉（列序、编码、日期写法），那条自救路径就是假的。
    """
    df = _sample()
    path = tmp_path / "exported.csv"
    path.write_bytes(export.to_csv_bytes(df))
    back = store.load_trades(path)

    assert ids(back) == ids(df)
    assert back["symbol"].tolist() == df["symbol"].tolist()   # 含 000333
    pd.testing.assert_frame_equal(back, df)


def test_empty_export_is_a_header_only_csv_not_a_zero_byte_file():
    """一笔都没筛到时也得给一份能打开的文件：只有表头。
    0 字节文件在 Excel 里是"文件已损坏"，而实情只是没有符合条件的记录。
    """
    blob = export.to_csv_bytes(store.empty_trades())
    assert blob[:3] == b"\xef\xbb\xbf"
    text = blob.decode("utf-8-sig")
    assert text.splitlines()[0] == ",".join(schema.COLUMNS)
    assert len(text.strip().splitlines()) == 1


def test_csv_writes_dates_as_iso_text():
    text = export.to_csv_bytes(_sample()).decode("utf-8-sig")
    assert "2026-08-03" in text
    assert "00:00:00" not in text, "日期不该带时间部分（datetime64 直接落盘的痕迹）"


# ================================================================ Excel 导出（§5.5）

def _sheet(blob: bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    return wb[export.SHEET_NAME]


def test_excel_can_be_reopened_by_openpyxl_with_the_same_header():
    sheet = _sheet(export.to_excel_bytes(_sample()))
    header = [c.value for c in sheet[1]]
    assert header == list(schema.COLUMNS)


def test_excel_keeps_chinese_correct():
    sheet = _sheet(export.to_excel_bytes(_sample()))
    values = [c.value for row in sheet.iter_rows() for c in row]
    assert "美的集团" in values and "20日线金叉" in values


def test_excel_writes_the_symbol_as_text_so_000333_stays_six_digits():
    """xlsx 里最容易丢的就是这个：数字单元格里的 000333 就是 333，
    而表格看上去完全正常。断言**类型**，不是"看着有三个零"。"""
    sheet = _sheet(export.to_excel_bytes(_sample()))
    col = list(schema.COLUMNS).index("symbol") + 1
    cell = sheet.cell(row=2, column=col)
    assert cell.value == "000333"
    assert isinstance(cell.value, str), f"symbol 被写成了 {type(cell.value)}"


def test_excel_writes_dates_as_iso_text_not_a_locale_dependent_datetime():
    """日期写成 datetime，打开时按机器的区域设置显示成 8/3/26 或 03.08.2026。
    ISO 文本在哪台机器上都是同一串，也与 CSV 那份逐字一致。"""
    sheet = _sheet(export.to_excel_bytes(_sample()))
    col = list(schema.COLUMNS).index("date") + 1
    cell = sheet.cell(row=2, column=col)
    assert cell.value == "2026-08-03"
    assert isinstance(cell.value, str), f"date 被写成了 {type(cell.value)}"


def test_excel_contains_exactly_the_filtered_rows():
    filtered = export.filter_trades(_sample(), sources=["discretionary"])
    sheet = _sheet(export.to_excel_bytes(filtered))
    col = list(schema.COLUMNS).index("trade_id") + 1
    got = [sheet.cell(row=r, column=col).value for r in range(2, sheet.max_row + 1)]
    assert got == ids(filtered) == ["T2", "T5"]


def test_empty_excel_still_has_the_header_row():
    sheet = _sheet(export.to_excel_bytes(store.empty_trades()))
    assert [c.value for c in sheet[1]] == list(schema.COLUMNS)
    assert sheet.max_row == 1


def test_missing_openpyxl_says_what_to_install(monkeypatch):
    """依赖没装时必须说清装什么。pandas 原生那句 "Missing optional dependency"
    不会提到本项目的装法，而这条路径只在别人 clone 之后第一次点导出时出现。"""
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    with pytest.raises(RuntimeError, match="openpyxl"):
        export.to_excel_bytes(_sample())


def test_export_filename_carries_the_day_and_the_extension():
    """文件名带日期：导出三次落到下载目录里，不带日期就是 trades(1).csv 这种谁也认不出的名字。"""
    assert export.export_name(date(2026, 8, 28), "csv") == "trades_20260828.csv"
    assert export.export_name(datetime(2026, 8, 28, 15, 4), "xlsx") == "trades_20260828.xlsx"


# ================================================================ 整表改删（§5.2 可编辑可删除）

def test_editing_one_field_leaves_every_other_row_byte_identical():
    df = _sample()
    edited = pd.DataFrame([{"trade_id": "T2", "reason": "改过的理由"}])
    out = store.apply_edits(df, edited)

    assert out.loc[out["trade_id"] == "T2", "reason"].item() == "改过的理由"
    pd.testing.assert_frame_equal(out.drop(index=1), df.drop(index=1))


def test_edits_are_located_by_trade_id_not_by_row_number():
    """编辑区显示的是**筛选后**的表：行号与整本日志对不上。

    按行号写盘的话，用户在筛出的第 2 行改一个字，改掉的是整本日志的第 2 行
    ——而且不会报错。设计 §2.3 明写 trade_id 是"编辑/删除靠它定位，不靠行号"。
    """
    df = _sample()
    out = store.apply_edits(df, pd.DataFrame([{"trade_id": "T5", "note": "已送到"}]))
    assert out.loc[out["trade_id"] == "T5", "note"].item() == "已送到"
    assert out.loc[out["trade_id"] == "T1", "note"].item() == ""


def test_unknown_trade_id_is_refused_instead_of_appended():
    """认不出的主键一律拒写。静默 append 会凭空多出一笔交易，
    而这份文件不可再生——多出来的那笔以后没人能判断真假。"""
    with pytest.raises(ValueError, match="trade_id"):
        store.apply_edits(_sample(), pd.DataFrame([{"trade_id": "T99", "reason": "x"}]))


def test_duplicate_trade_id_in_the_edit_is_refused():
    """同一个主键给两套值，哪个赢都是猜的。"""
    edited = pd.DataFrame([{"trade_id": "T1", "reason": "甲"},
                           {"trade_id": "T1", "reason": "乙"}])
    with pytest.raises(ValueError, match="重复"):
        store.apply_edits(_sample(), edited)


def test_blank_trade_id_is_refused():
    """空主键定位不到任何一行。放行等于"改了但什么都没改"，最难发现。"""
    with pytest.raises(ValueError, match="trade_id"):
        store.apply_edits(_sample(), pd.DataFrame([{"trade_id": "", "reason": "x"}]))


def test_unknown_column_in_the_edit_is_refused():
    with pytest.raises(ValueError, match="未知字段"):
        store.apply_edits(_sample(), pd.DataFrame([{"trade_id": "T1", "symbol_code": "1"}]))


def test_delete_removes_exactly_the_named_rows():
    out = store.apply_edits(_sample(), pd.DataFrame(), delete_ids=["T2", "T4"])
    assert ids(out) == ["T1", "T3", "T5"]


def test_delete_of_an_unknown_id_is_refused():
    """"删掉了但其实没删"是最坏的结果：用户以为那笔已经不算了，盈亏还在算它。"""
    with pytest.raises(ValueError, match="trade_id"):
        store.apply_edits(_sample(), pd.DataFrame(), delete_ids=["T99"])


def test_row_order_is_preserved_so_the_diff_against_the_backup_stays_small():
    """行序稳定 = 改动看得出来。日志自 v0.3.2 起不在版本控制里（它是用户数据），
    但对比的对象只是从"git 上一次提交"换成了 `trades.csv.bak`（上一版备份）：
    每次保存都重排行序的话，那个对比是整个文件，等于看不出到底改了什么。"""
    df = _sample()
    out = store.apply_edits(df, pd.DataFrame([{"trade_id": "T4", "reason": "止损"}]))
    assert ids(out) == ids(df)


def test_edit_and_delete_in_one_call():
    out = store.apply_edits(_sample(),
                            pd.DataFrame([{"trade_id": "T1", "reason": "改了"}]),
                            delete_ids=["T5"])
    assert ids(out) == ["T1", "T2", "T3", "T4"]
    assert out.loc[0, "reason"] == "改了"


def test_edits_keep_dtypes_and_leading_zeros():
    out = store.apply_edits(_sample(),
                            pd.DataFrame([{"trade_id": "T1", "shares": "1200"}]))
    assert out["shares"].dtype == "float64"
    assert out.loc[0, "shares"] == 1200.0
    assert out.loc[0, "symbol"] == "000333"
    assert pd.api.types.is_datetime64_any_dtype(out["date"])


def test_apply_edits_does_not_touch_the_input_frames():
    df = _sample()
    before = df.copy(deep=True)
    edited = pd.DataFrame([{"trade_id": "T1", "reason": "改了"}])
    edited_before = edited.copy(deep=True)
    store.apply_edits(df, edited, delete_ids=["T5"])
    pd.testing.assert_frame_equal(df, before)
    pd.testing.assert_frame_equal(edited, edited_before)


def test_editing_the_date_is_accepted_as_iso_text():
    """编辑器交回来的日期可能是字符串（用户手打）。认不出来必须抛，不许静默变 NaT
    ——那笔交易会从所有按日期筛选的视图里消失，而汇总数字照样给得出来。"""
    out = store.apply_edits(_sample(),
                            pd.DataFrame([{"trade_id": "T1", "date": "2026-08-04"}]))
    assert out.loc[0, "date"] == pd.Timestamp("2026-08-04")
    with pytest.raises(Exception):
        store.apply_edits(_sample(),
                          pd.DataFrame([{"trade_id": "T1", "date": "八月四日"}]))


def test_apply_edits_of_nothing_is_a_no_op():
    df = _sample()
    pd.testing.assert_frame_equal(store.apply_edits(df, pd.DataFrame()), df)


def test_saving_the_edited_table_survives_a_round_trip(tmp_path):
    """改完 → 落盘 → 读回，还得是同一张表（这是页面上"保存修改"那条路的全程）。"""
    path = tmp_path / "trades.csv"
    store.save_trades(_sample(), path)
    out = store.apply_edits(store.load_trades(path),
                            pd.DataFrame([{"trade_id": "T3", "note": "已到账"}]),
                            delete_ids=["T2"])
    store.save_trades(out, path)
    back = store.load_trades(path)

    assert ids(back) == ["T1", "T3", "T4", "T5"]
    assert back.loc[1, "note"] == "已到账"
    assert back["symbol"].tolist() == ["000333", "000333", "000333", "600519"]


# ================================================================ 落盘与导出共用一份序列化

def test_serialize_turns_dates_into_iso_text_and_keeps_symbols_as_text():
    out = store.serialize(_sample())
    assert out.loc[0, "date"] == "2026-08-03"
    assert out.loc[0, "symbol"] == "000333"


def test_serialize_does_not_touch_the_input_frame():
    df = _sample()
    before = df.copy(deep=True)
    store.serialize(df)
    pd.testing.assert_frame_equal(df, before)
